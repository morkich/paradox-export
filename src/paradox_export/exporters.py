"""Converters from decoded table rows to CSV, XML and XLSX files."""

import csv
import datetime
import re
from xml.sax.saxutils import escape

from .reader import PdxTable

DATE_RE = re.compile(r'^(\d{2})\.(\d{2})\.(\d{4})$')
TIME_RE = re.compile(r'^(\d{1,2}):(\d{2})(?::(\d{2}))?$')


def prepare(path, unique=False):
    """Read a table and return (header, rows).

    Duplicates of the first (numeric) field — a known artefact of files that
    went through abnormal shutdowns or manual copying:

    * ``unique=False`` (default): keep all rows and, if duplicates exist,
      prepend a ``_dup_key`` column marking them.
    * ``unique=True``: keep only the last physical copy of every key.
    """
    t = PdxTable(path)
    header = list(t.field_names)
    rows = list(t.rows())
    can_key = bool(t.field_types) and t.field_types[0] in (3, 4, 0x16)
    if not can_key:
        return header, rows
    seen = {}
    for r in rows:
        if r[0]:
            seen[r[0]] = seen.get(r[0], 0) + 1
    dups = {k for k, v in seen.items() if v > 1}
    if not dups:
        return header, rows
    if unique:
        last = {}
        for i, r in enumerate(rows):
            if r[0]:
                last[r[0]] = i
        return header, [r for i, r in enumerate(rows) if not r[0] or last.get(r[0]) == i]
    return (['_dup_key'] + header), [[('Y' if r[0] in dups else '')] + r for r in rows]


def export_csv(header, rows, out_path):
    """Write a UTF-8 (BOM) CSV with ';' as the separator (Excel friendly)."""
    with open(out_path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';', quoting=csv.QUOTE_MINIMAL)
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


def export_xml(name, header, rows, out_path):
    """Write a simple XML document: <name><row><field>value</field>...</row>..."""
    n = 0
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('<?xml version="1.0" encoding="utf-8"?>\n')
        f.write('<%s>\n' % name)
        f.write('  <fields>' + ''.join('<field name="%s"/>' % escape(h) for h in header) + '</fields>\n')
        for r in rows:
            parts = []
            for h, v in zip(header, r):
                if v in ('', None):
                    continue
                parts.append('<%s>%s</%s>' % (h, escape(str(v)), h))
            f.write('  <row>' + ''.join(parts) + '</row>\n')
            n += 1
        f.write('</%s>\n' % name)
    return n


def xlsx_available():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def _parse_date(s):
    m = DATE_RE.match(s) if isinstance(s, str) else None
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        return datetime.datetime(y, mo, d)
    except ValueError:
        return None


def _parse_time(s):
    m = TIME_RE.match(s.strip()) if isinstance(s, str) else None
    if not m:
        return None
    h, mi, sec = int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)
    if h > 23 or mi > 59 or sec > 59:
        return None
    return datetime.time(h, mi, sec)


def _parse_num(s):
    if not isinstance(s, str):
        return None
    t = s.strip()
    if len(t) > 1 and t[0] == '0' and t[1].isdigit():
        return None  # leading zeros (zip codes etc) -> keep as text
    try:
        return float(t)
    except ValueError:
        return None


def _detect_modes(header, rows):
    """Guess the column type from a sample: 'date' / 'time' / 'num' / 'text'."""
    modes = []
    sample = rows[:500]
    for c in range(len(header)):
        vals = [r[c] for r in sample if r[c] not in ('', None)]
        if not vals:
            modes.append('text')
        elif all(_parse_date(v) is not None for v in vals):
            modes.append('date')
        elif all(_parse_time(v) is not None for v in vals):
            modes.append('time')
        elif all(_parse_num(v) is not None for v in vals):
            modes.append('num')
        else:
            modes.append('text')
    return modes


def export_xlsx(tables, out_path):
    """Write all tables into one workbook.

    ``tables`` is a list of ``(sheet_name, header, rows)``.
    Returns the total number of rows written.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='4472C4')
    used, total = set(), 0
    for base_name, header, rows in tables:
        name = re.sub(r'[\\/*?:\[\]]', '_', base_name)[:31] or 'Sheet'
        cand, i = name, 1
        while cand in used:
            cand = (name[:29] + '~%d' % i)[:31]
            i += 1
        used.add(cand)
        ws = wb.create_sheet(cand)
        ws.append(header)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
        modes = _detect_modes(header, rows)
        conv_rows = []
        for r in rows:
            out = []
            for c, v in enumerate(r):
                if v in ('', None):
                    out.append(None)
                elif modes[c] == 'date':
                    out.append(_parse_date(v) or v)
                elif modes[c] == 'time':
                    out.append(_parse_time(v) or v)
                elif modes[c] == 'num':
                    n = _parse_num(v)
                    out.append(v if n is None else n)
                else:
                    out.append(v)
            conv_rows.append(out)
        for rr in conv_rows:
            ws.append(rr)
        for c, m in enumerate(modes):
            if m in ('date', 'time'):
                fmt = 'DD.MM.YYYY' if m == 'date' else 'HH:MM:SS'
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=c + 1).number_format = fmt
        widths = [len(h) for h in header]
        for r in rows[:300]:
            for c, v in enumerate(r):
                if c < len(widths):
                    widths[c] = max(widths[c], len(str(v)))
        for c, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(c + 1)].width = min(w + 2, 60)
        ws.freeze_panes = 'A2'
        if header:
            ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(header)), max(ws.max_row, 1))
        total += len(rows)
    wb.save(out_path)
    return total
